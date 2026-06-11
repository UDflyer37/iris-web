#  OhCR / DCOE report export — multiple formats for NETO / observers

from __future__ import annotations

import datetime
import os
import re
import tempfile

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt, RGBColor

from app.datamgmt.dcoe.dcoe_report_forms_db import (
    _format_field_value,
    _section_instances,
    compile_form_to_markdown,
    get_report_form_schema,
)

EXPORT_FORMATS = {
    "docx": {
        "label": "Word (.docx)",
        "hint": "Best for NETO — open in Microsoft Word",
        "extension": ".docx",
        "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    },
    "html": {
        "label": "Web page (.html)",
        "hint": "Opens in any browser — print to PDF from File → Print",
        "extension": ".html",
        "mime": "text/html",
    },
    "md": {
        "label": "Markdown (.md)",
        "hint": "Plain structured text for technical teams",
        "extension": ".md",
        "mime": "text/markdown",
    },
    "txt": {
        "label": "Plain text (.txt)",
        "hint": "Simple text for email or chat",
        "extension": ".txt",
        "mime": "text/plain",
    },
}


def list_export_formats() -> list[dict]:
    return [{"id": key, **meta} for key, meta in EXPORT_FORMATS.items()]


def _base_filename(draft: dict, case_name: str) -> str:
    safe_case = re.sub(r"[^\w\-]+", "_", case_name or "case")[:40]
    template_key = draft.get("report_template_key", "report")
    date_stamp = (draft.get("updated_at") or draft.get("created_at") or "")[:10] or "export"
    return f"OhCR-{safe_case}-{template_key}-{date_stamp}"


def _report_context(draft: dict) -> tuple[dict, dict]:
    schema = get_report_form_schema(draft.get("form_schema_key") or draft.get("report_template_key"))
    form_data = dict(draft.get("form_data") or {})
    if not form_data.get("_report_date"):
        form_data["_report_date"] = datetime.datetime.utcnow().strftime("%Y-%m-%d")
    return schema, form_data


def _metadata_rows(schema: dict, form_data: dict) -> list[tuple[str, str]]:
    rows = [
        ("TLP", schema.get("tlp") or "AMBER+STRICT"),
        ("Date", form_data.get("_report_date", "—")),
        ("Case", form_data.get("_case_name") or "—"),
        ("Author", form_data.get("_author") or "—"),
    ]
    customer = (form_data.get("_customer") or "").strip()
    if customer:
        rows.append(("Customer", customer))
    return rows


def _add_docx_field(doc: Document, field: dict, value) -> None:
    label = field.get("label", field["id"])
    display = _format_field_value(field, value)
    field_type = field.get("type", "text")

    if field_type == "textarea":
        p = doc.add_paragraph()
        run = p.add_run(f"{label}")
        run.bold = True
        run.font.size = Pt(11)
        text = str(value or "").strip()
        body = doc.add_paragraph(text if text else "Not provided")
        body.paragraph_format.space_after = Pt(8)
        return

    p = doc.add_paragraph()
    label_run = p.add_run(f"{label}: ")
    label_run.bold = True
    label_run.font.size = Pt(11)
    value_run = p.add_run(display)
    value_run.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(4)


def _build_docx(draft: dict) -> Document:
    schema, form_data = _report_context(draft)
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = schema.get("title") or draft.get("title") or "OhCR Report"
    doc.add_heading(title, level=0)

    meta_table = doc.add_table(rows=len(_metadata_rows(schema, form_data)), cols=2)
    meta_table.style = "Table Grid"
    for idx, (label, value) in enumerate(_metadata_rows(schema, form_data)):
        meta_table.rows[idx].cells[0].text = label
        meta_table.rows[idx].cells[0].paragraphs[0].runs[0].bold = True
        meta_table.rows[idx].cells[1].text = value
    doc.add_paragraph()

    if schema.get("warning"):
        warn = doc.add_paragraph(schema["warning"])
        warn.runs[0].italic = True
        warn.runs[0].font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
        doc.add_paragraph()

    for section in schema.get("sections", []):
        if section.get("repeatable"):
            instances = _section_instances(form_data, section)
            base_label = section.get("label", "Section")
            for inst_idx, instance in enumerate(instances, 1):
                heading = f"{base_label} — Entry {inst_idx}" if len(instances) > 1 else base_label
                doc.add_heading(heading, level=1)
                for field in section.get("fields", []):
                    _add_docx_field(doc, field, instance.get(field["id"]))
        else:
            doc.add_heading(section.get("label", "Section"), level=1)
            for field in section.get("fields", []):
                _add_docx_field(doc, field, form_data.get(field["id"]))

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = f"Generated {datetime.datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')} — OhCR/DCOE"
    footer.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(0x75, 0x75, 0x75)

    return doc


def _build_plain_text(draft: dict) -> str:
    schema, form_data = _report_context(draft)
    lines: list[str] = []
    title = schema.get("title") or draft.get("title") or "OhCR Report"

    lines.append(title.upper())
    lines.append("=" * len(title))
    lines.append("")

    for label, value in _metadata_rows(schema, form_data):
        lines.append(f"{label}: {value}")
    lines.append("")

    if schema.get("warning"):
        lines.append(f"NOTICE: {schema['warning']}")
        lines.append("")

    for section in schema.get("sections", []):
        if section.get("repeatable"):
            instances = _section_instances(form_data, section)
            base_label = section.get("label", "Section")
            for inst_idx, instance in enumerate(instances, 1):
                heading = f"{base_label} — Entry {inst_idx}" if len(instances) > 1 else base_label
                lines.append(heading)
                lines.append("-" * len(heading))
                for field in section.get("fields", []):
                    label = field.get("label", field["id"])
                    value = _format_field_value(field, instance.get(field["id"]))
                    if field.get("type") == "textarea":
                        lines.append(f"{label}:")
                        text = str(instance.get(field["id"]) or "").strip()
                        lines.append(text if text else "Not provided")
                    else:
                        lines.append(f"{label}: {value}")
                lines.append("")
        else:
            heading = section.get("label", "Section")
            lines.append(heading)
            lines.append("-" * len(heading))
            for field in section.get("fields", []):
                label = field.get("label", field["id"])
                value = _format_field_value(field, form_data.get(field["id"]))
                if field.get("type") == "textarea":
                    lines.append(f"{label}:")
                    text = str(form_data.get(field["id"]) or "").strip()
                    lines.append(text if text else "Not provided")
                else:
                    lines.append(f"{label}: {value}")
            lines.append("")

    lines.append(f"Generated {datetime.datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")
    return "\n".join(lines).strip() + "\n"


def _escape_html(text: str) -> str:
    return (
        str(text or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _build_html(draft: dict) -> str:
    schema, form_data = _report_context(draft)
    title = schema.get("title") or draft.get("title") or "OhCR Report"
    tlp = schema.get("tlp") or "AMBER+STRICT"

    body_parts = [f'<h1 class="rpt-title">{_escape_html(title)}</h1>']
    body_parts.append('<table class="rpt-meta">')
    for label, value in _metadata_rows(schema, form_data):
        body_parts.append(
            f"<tr><th>{_escape_html(label)}</th><td>{_escape_html(value)}</td></tr>"
        )
    body_parts.append("</table>")

    if schema.get("warning"):
        body_parts.append(f'<div class="rpt-notice">{_escape_html(schema["warning"])}</div>')

    for section in schema.get("sections", []):
        if section.get("repeatable"):
            instances = _section_instances(form_data, section)
            base_label = section.get("label", "Section")
            for inst_idx, instance in enumerate(instances, 1):
                heading = f"{base_label} — Entry {inst_idx}" if len(instances) > 1 else base_label
                body_parts.append(f'<h2 class="rpt-section">{_escape_html(heading)}</h2>')
                for field in section.get("fields", []):
                    label = field.get("label", field["id"])
                    if field.get("type") == "textarea":
                        text = str(instance.get(field["id"]) or "").strip() or "Not provided"
                        body_parts.append(f'<p class="rpt-label">{_escape_html(label)}</p>')
                        body_parts.append(f'<p class="rpt-body">{_escape_html(text).replace(chr(10), "<br>")}</p>')
                    else:
                        val = _format_field_value(field, instance.get(field["id"]))
                        body_parts.append(
                            f'<p class="rpt-field"><strong>{_escape_html(label)}:</strong> {_escape_html(val)}</p>'
                        )
        else:
            body_parts.append(f'<h2 class="rpt-section">{_escape_html(section.get("label", "Section"))}</h2>')
            for field in section.get("fields", []):
                label = field.get("label", field["id"])
                if field.get("type") == "textarea":
                    text = str(form_data.get(field["id"]) or "").strip() or "Not provided"
                    body_parts.append(f'<p class="rpt-label">{_escape_html(label)}</p>')
                    body_parts.append(f'<p class="rpt-body">{_escape_html(text).replace(chr(10), "<br>")}</p>')
                else:
                    val = _format_field_value(field, form_data.get(field["id"]))
                    body_parts.append(
                        f'<p class="rpt-field"><strong>{_escape_html(label)}:</strong> {_escape_html(val)}</p>'
                    )

    body = "\n".join(body_parts)
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <title>{_escape_html(title)}</title>
  <style>
    body {{ font-family: Calibri, 'Segoe UI', Arial, sans-serif; max-width: 800px; margin: 2em auto; padding: 0 1.5em; color: #222; line-height: 1.55; }}
    .tlp-banner {{ background: #fff3cd; border: 1px solid #ffc107; padding: 10px 14px; margin-bottom: 1.5em; font-weight: bold; }}
    .rpt-title {{ color: #1a237e; border-bottom: 2px solid #c5cae9; padding-bottom: 8px; }}
    .rpt-meta {{ width: 100%; border-collapse: collapse; margin: 1em 0 1.5em; }}
    .rpt-meta th {{ text-align: left; width: 120px; padding: 6px 10px; background: #f5f5f5; border: 1px solid #ddd; }}
    .rpt-meta td {{ padding: 6px 10px; border: 1px solid #ddd; }}
    .rpt-section {{ color: #37474f; margin-top: 1.5em; font-size: 1.15em; }}
    .rpt-label {{ font-weight: bold; margin: 0.8em 0 0.2em; color: #455a64; }}
    .rpt-body {{ margin: 0 0 1em; white-space: pre-wrap; }}
    .rpt-field {{ margin: 0.3em 0; }}
    .rpt-notice {{ background: #fff3cd; border-left: 4px solid #ffc107; padding: 10px 14px; margin: 1em 0; }}
    .no-print {{ font-size: 12px; color: #666; }}
    @media print {{ body {{ margin: 0; }} .no-print {{ display: none; }} }}
  </style>
</head>
<body>
  <div class="tlp-banner">TLP: {_escape_html(tlp)}</div>
  <p class="no-print">To save as PDF: File → Print → Save as PDF</p>
  {body}
</body>
</html>"""


def draft_export_path(draft: dict, case_name: str, export_format: str = "docx") -> tuple[str, str, str, str]:
    """Returns (file_path, download_filename, mime_type, temp_directory)."""
    fmt = (export_format or "docx").lower()
    if fmt not in EXPORT_FORMATS:
        raise ValueError(f"Unsupported export format: {export_format}")

    meta = EXPORT_FORMATS[fmt]
    base = _base_filename(draft, case_name)
    filename = f"{base}{meta['extension']}"

    tmp_dir = tempfile.mkdtemp(prefix="dcoe_export_")
    fpath = os.path.join(tmp_dir, filename)

    if fmt == "md":
        schema, form_data = _report_context(draft)
        content = draft.get("content") or compile_form_to_markdown(schema, form_data)
        with open(fpath, "w", encoding="utf-8") as handle:
            handle.write(content)
    elif fmt == "txt":
        with open(fpath, "w", encoding="utf-8") as handle:
            handle.write(_build_plain_text(draft))
    elif fmt == "html":
        with open(fpath, "w", encoding="utf-8") as handle:
            handle.write(_build_html(draft))
    elif fmt == "docx":
        document = _build_docx(draft)
        document.save(fpath)

    if not os.path.isfile(fpath):
        raise OSError(f"Export file was not created: {fpath}")

    return fpath, filename, meta["mime"], tmp_dir
